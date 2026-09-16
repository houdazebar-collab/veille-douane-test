import re, copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from lxml import etree

NAVY='004079'; NAVY2='063A68'; TEAL='00A49F'; SKY='0092D2'; RED='D8232A'; PINK='DF337F'; ORANGE='F18700'
GREY='5F6B76'; INK='1A2733'; LINE='E1E9F0'; PALE='F4F7FA'; ICE='E6F0F8'; AMBER='FFF6E5'; GREEN='1E6B3A'
SRC='AtelierRDUECCI_revu_2026-09-16.pptx'; OUT='AtelierRDUECCI_revu_v2.pptx'
p=Presentation(SRC)
NS='{http://schemas.openxmlformats.org/drawingml/2006/main}'

def rgb(h): return RGBColor.from_string(h)
def shadow(shape):
    spPr=shape._element.spPr
    eff=etree.SubElement(spPr, NS+'effectLst')
    sh=etree.SubElement(eff, NS+'outerShdw', blurRad="76200", rotWithShape="0", algn="bl", dir="5400000", dist="25400")
    c=etree.SubElement(sh, NS+'srgbClr', val="8AA0B4"); etree.SubElement(c, NS+'alpha', val="27843")
def box(slide,x,y,w,h,fill=None,line=None,radius=None,shape=MSO_SHAPE.RECTANGLE,shdw=False,dash=False,lw=1.0):
    s=slide.shapes.add_shape(shape,Inches(x),Inches(y),Inches(w),Inches(h))
    if radius is not None and shape==MSO_SHAPE.ROUNDED_RECTANGLE: s.adjustments[0]=radius
    if fill: s.fill.solid(); s.fill.fore_color.rgb=rgb(fill)
    else: s.fill.background()
    if line: s.line.color.rgb=rgb(line); s.line.width=Pt(lw)
    else: s.line.fill.background()
    if dash:
        ln=s._element.spPr.find(NS+'ln'); etree.SubElement(ln, NS+'prstDash', val='dash')
    if shdw: shadow(s)
    s.text_frame.text=''  # keep empty
    return s
def text(slide,x,y,w,h,paras,size=11,color=INK,bold=False,align='l',anchor='ctr',margin=0.06,italic=False,bullets=False,spc=0):
    """paras: list of str or list of (str, {size,color,bold}) runs lists"""
    t=slide.shapes.add_textbox(Inches(x),Inches(y),Inches(w),Inches(h)); tf=t.text_frame; tf.word_wrap=True
    tf.margin_left=tf.margin_right=Inches(margin); tf.margin_top=tf.margin_bottom=Inches(0.03)
    tf.vertical_anchor={'ctr':MSO_ANCHOR.MIDDLE,'t':MSO_ANCHOR.TOP,'b':MSO_ANCHOR.BOTTOM}[anchor]
    if isinstance(paras,str): paras=[paras]
    for i,para in enumerate(paras):
        par=tf.paragraphs[0] if i==0 else tf.add_paragraph()
        par.alignment={'l':PP_ALIGN.LEFT,'c':PP_ALIGN.CENTER,'r':PP_ALIGN.RIGHT}[align]
        if spc: par.space_after=Pt(spc)
        runs=[(para,{})] if isinstance(para,str) else para
        for txt,o in runs:
            r=par.add_run(); r.text=txt; f=r.font; f.name='Calibri'; f.size=Pt(o.get('size',size)); f.bold=o.get('bold',bold); f.italic=o.get('italic',italic)
            f.color.rgb=rgb(o.get('color',color))
        if bullets:
            pPr=par._p.get_or_add_pPr(); pPr.set('marL','171450'); pPr.set('indent','-171450')
            bc=etree.SubElement(pPr,NS+'buClr'); etree.SubElement(bc,NS+'srgbClr',val=NAVY)
            etree.SubElement(pPr,NS+'buChar',char='•')
    return t
def pill(slide,x,y,w,h,fill,label,size=9.5,color='FFFFFF',bold=True):
    box(slide,x,y,w,h,fill=fill,radius=0.5,shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(slide,x,y,w,h,label,size=size,color=color,bold=bold,align='c')
def arrow_down(slide,x,y,w=0.28,h=0.22,color=NAVY):
    s=slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,Inches(x),Inches(y),Inches(w),Inches(h)); s.rotation=180
    s.fill.solid(); s.fill.fore_color.rgb=rgb(color); s.line.fill.background(); return s
def arrow_right(slide,x,y,w=0.3,h=0.3,color=NAVY):
    s=slide.shapes.add_shape(MSO_SHAPE.RIGHT_ARROW,Inches(x),Inches(y),Inches(w),Inches(h))
    s.fill.solid(); s.fill.fore_color.rgb=rgb(color); s.line.fill.background(); return s
def clear_pictures(slide):
    for sh in list(slide.shapes):
        if sh.shape_type==13 and Emu(sh.width).inches>5: sh._element.getparent().remove(sh._element)
def set_title(slide,title,sub):
    text(slide,0.5,0.42,7.2,0.70,title,size=27,color=NAVY,bold=True,margin=0.13)
    text(slide,0.5,1.08,8.9,0.40,sub,size=13,color=GREY,margin=0.13)

# ================= SLIDE 12 : entrant / sortant =================
s=p.slides[11]; clear_pictures(s)
set_title(s,"Qui est concerné ? Entrant et sortant","Une fois le produit dans le champ (annexe I + matière) : que se passe-t-il selon le flux ?")
LX,LW=0.5,4.05; RX,RW=5.45,4.05; AX=4.62
pill(s,LX+0.6,1.55,LW-1.2,0.36,NAVY,"PRODUIT ENTRANT",size=11)
pill(s,RX+0.6,1.55,RW-1.2,0.36,NAVY,"PRODUIT SORTANT",size=11)
rows=[
 (("Importation sur le marché de l'UE",[("Vous = opérateur",{'bold':True}),(" · DDR + n° en douane (C716)",{})]),(TEAL,"DDR requise"),
  ("Exportation depuis le marché de l'UE",[("Vous = opérateur",{'bold':True}),(" · DDR avant l'export",{})]),(TEAL,"DDR requise"),True),
 (("Acheté sur le marché UE et utilisé pour un produit RDUE sortant",[("Opérateur en aval",{'bold':True}),(" · pas de DDR, conserve le n° du fournisseur",{})]),(NAVY,"N° de DDR à conserver"),
  ("Vendu au sein du marché UE (« mise à disposition »)",[("Commerçant",{'bold':True}),(" · références fournisseur conservées",{})]),(NAVY,"N° transmis par le fournisseur"),True),
 (("Acheté sur le marché UE et PAS utilisé pour un produit RDUE sortant",[("Produit non relevant",{'bold':True,'color':RED})]),(RED,"Pas de DDR"),
  ("Exporté ou vendu sur le marché de l'UE",[("Produit non relevant",{'bold':True,'color':RED})]),(RED,"Pas de DDR"),True),
 (("Acheté sur le marché UE — consommation interne",[("Produit non relevant",{'bold':True,'color':RED})]),(RED,"Pas de DDR"),
  ("Hors obligation",[("Aucune DDR, aucun n° à conserver",{'bold':True,'color':RED})]),(RED,"Pas de DDR"),False),
]
y=2.02; RH=1.05; GAP=0.12
for (lt,ld),(lc,ll),(rt,rd),(rc,rl),arr in rows:
    for X,W,(t1,d1),(c1,l1) in ((LX,LW,(lt,ld),(lc,ll)),(RX,RW,(rt,rd),(rc,rl))):
        box(s,X,y,W,RH,fill='FFFFFF',line=(RED if c1==TEAL else LINE),radius=0.08,shape=MSO_SHAPE.ROUNDED_RECTANGLE,shdw=True,lw=(1.5 if c1==TEAL else 1))
        text(s,X+0.12,y+0.05,W-0.24,0.5,t1,size=11,color=NAVY,bold=True,anchor='t')
        text(s,X+0.12,y+0.52,W-0.24,0.24,[d1],size=9,color=GREY,anchor='t')
        pill(s,X+0.15,y+RH-0.32,2.1,0.24,c1,l1,size=8.5)
    if arr: arrow_right(s,AX,y+RH/2-0.15,0.32,0.3,color=NAVY)
    else: text(s,AX,y+RH/2-0.2,0.32,0.4,"✕",size=16,color=GREY,bold=True,align='c')
    y+=RH+GAP
box(s,0.5,y+0.0,9.0,0.46,fill=AMBER,line='F3D9A4',radius=0.15,shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s,0.6,y+0.0,8.8,0.46,[[("Régimes douaniers particuliers : ",{'bold':True,'color':NAVY}),("entrepôt douanier, perfectionnement actif, admission temporaire… ne sont pas soumis au RDUE (FAQ de la Commission).",{})]],size=9.5,color=INK)
s.notes_slide.notes_text_frame.text="Produit entrant → produit sortant. Faire répondre la salle. Depuis 2025/2650, l'aval ne dépose pas de DDR : il conserve le n° de son fournisseur."

# ================= SLIDE 20 : logigramme =================
s=p.slides[19]; clear_pictures(s)
set_title(s,"Le parcours de diligence raisonnée","Le parcours obligatoire, article par article (art. 8 à 13 du règlement).")
# Étape 1
box(s,0.5,1.6,9.0,1.15,fill='FFFFFF',line=LINE,radius=0.06,shape=MSO_SHAPE.ROUNDED_RECTANGLE,shdw=True)
pill(s,0.65,1.72,0.36,0.36,NAVY,"1",size=12)
text(s,1.1,1.68,5,0.42,"Collecte des informations",size=13,color=NAVY,bold=True)
pill(s,8.3,1.75,1.05,0.28,RED,"ARTICLE 9",size=8.5)
text(s,0.75,2.1,4.2,0.62,["Description, nom commercial, type de produit","Quantités et pays (zones) de production","Fournisseurs et clients (nom, adresse)"],size=9.5,anchor='t',bullets=True)
text(s,5.0,2.1,4.4,0.62,[[("Géolocalisation de toutes les parcelles",{'bold':True})],"Date ou période de production","Preuves de légalité et de « zéro déforestation »"],size=9.5,anchor='t',bullets=True)
arrow_down(s,6.7,2.82)
# Décision
box(s,4.2,3.1,5.3,0.62,fill=GREEN,radius=0.15,shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s,4.3,3.1,5.1,0.62,[[("Pays de production à risque standard ou élevé ?",{'bold':True,'size':11.5,'color':'FFFFFF'})],[("classification des pays (benchmarking, art. 29)",{'size':9,'color':'D7EBDD'})]],align='c')
# Branche risque faible
text(s,3.72,3.2,0.5,0.42,"NON",size=9.5,color=GREEN,bold=True,align='c')
s_=s.shapes.add_shape(MSO_SHAPE.LEFT_ARROW,Inches(3.85),Inches(3.55),Inches(0.32),Inches(0.26)); s_.fill.solid(); s_.fill.fore_color.rgb=rgb(GREEN); s_.line.fill.background()
box(s,0.5,3.1,3.3,1.95,fill=ICE,line=SKY,radius=0.06,shape=MSO_SHAPE.ROUNDED_RECTANGLE,dash=True)
pill(s,0.65,3.22,1.9,0.26,'FFFFFF',"pays à risque faible",size=8.5,color=NAVY)
text(s,0.65,3.55,3.0,0.4,"Diligence raisonnée simplifiée",size=12,color=NAVY,bold=True,anchor='t')
text(s,0.65,3.92,3.0,0.9,[[("Seule l'étape 1 est exigée. ",{'bold':True}),("L'opérateur n'est pas tenu à l'évaluation ni à l'atténuation du risque (art. 10 et 11).",{})]],size=9.5,anchor='t')
pill(s,2.55,4.68,1.1,0.26,RED,"ARTICLE 13",size=8.5)
text(s,4.2,3.75,0.5,0.3,"OUI",size=9.5,color=GREEN,bold=True,align='c')
arrow_down(s,6.7,3.78,color=GREEN)
# Étapes 2 et 3
for X,num,ttl,art,items in ((4.2,"2","Évaluation du risque","ARTICLE 10",["Vérifier et analyser les informations (critères de l'art. 10)","Réexamen au moins une fois par an","Pas de mise sur le marché si le risque n'est pas nul ou négligeable"]),
                            (6.95,"3","Atténuation du risque","ARTICLE 11",["Informations, données ou documents complémentaires","Audits ou enquêtes indépendants","Procédures et contrôles proportionnés"])):
    box(s,X,4.05,2.55,1.55,fill='FFFFFF',line=(SKY if num=="2" else ORANGE),radius=0.06,shape=MSO_SHAPE.ROUNDED_RECTANGLE,shdw=True,lw=1.5)
    pill(s,X+0.12,4.15,0.3,0.3,(SKY if num=="2" else ORANGE),num,size=10)
    text(s,X+0.48,4.12,2.0,0.36,ttl,size=11.5,color=NAVY,bold=True)
    text(s,X+0.1,4.5,2.35,0.85,items,size=8.5,anchor='t',bullets=True)
    pill(s,X+1.45,5.3,1.0,0.24,RED,art,size=8)
# Résultats
box(s,4.2,5.75,2.55,0.42,fill='E3F5F3',line=TEAL,radius=0.3,shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s,4.2,5.75,2.55,0.42,"Risque nul ou négligeable",size=10,color=TEAL,bold=True,align='c')
box(s,6.95,5.75,2.55,0.42,fill='FDECEC',line=RED,radius=0.3,shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s,6.95,5.75,2.55,0.42,"Risque non négligeable : STOP (art. 3)",size=10,color=RED,bold=True,align='c')
arrow_down(s,5.34,6.2,color=TEAL)
ln=s.shapes.add_connector(1,Inches(2.14),Inches(5.05),Inches(2.14),Inches(6.1)); ln.line.color.rgb=rgb(SKY); ln.line.width=Pt(2.25)
arrow_down(s,2.0,6.1,color=SKY)
# DDR
box(s,0.5,6.35,9.0,0.62,fill=NAVY,radius=0.08,shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s,0.65,6.35,7.5,0.62,[[("Déclaration de diligence raisonnée",{'bold':True,'size':12.5,'color':'FFFFFF'})],[("Déposée dans le système d'information avant la mise sur le marché ou l'export → n° de référence en douane (C716) et au 1er client aval.",{'size':9,'color':'CFE0EE'})]])
pill(s,8.25,6.52,1.1,0.28,'FFFFFF',"ART. 4 & 12",size=8.5,color=NAVY)
s.notes_slide.notes_text_frame.text="Logigramme diligence (art. 8-13). Cœur de la conformité. Pays à risque faible = étape 1 seulement."

# ================= Titres trop longs =================
def retitle(n,new):
    for sh in p.slides[n-1].shapes:
        if sh.has_text_frame and any(r.font.size and r.font.size.pt==27 for par in sh.text_frame.paragraphs for r in par.runs):
            sh.text_frame.paragraphs[0].runs[0].text=new
            for r in sh.text_frame.paragraphs[0].runs[1:]: r.text=''
            return
retitle(10,"Annexe I : ce qui change en 2026-2027")
retitle(11,"Trois conditions cumulatives (art. 3)")
retitle(14,"Cas pratiques : quel rôle pour qui ?")
retitle(15,"Précisions clés — FAQ Commission v5")
retitle(19,"Diligence simplifiée : le rôle du pays")
for sh in p.slides[18].shapes:
    if sh.has_text_frame and sh.text_frame.text.startswith("Le pays de production module"):
        sh.text_frame.paragraphs[0].runs[0].text="Le pays de production module la diligence (art. 29) — liste du 22/05/2025, révision annoncée."

# ================= Plancher de taille de police =================
CHROME=re.compile(r'^(RDUE / EUDR|\d+ / \d+$|CCI|Saint-Étienne)')
for i,sl in enumerate(p.slides,1):
    if i in (12,20): continue
    for sh in sl.shapes:
        if not sh.has_text_frame or CHROME.match(sh.text_frame.text.strip()): continue
        for par in sh.text_frame.paragraphs:
            for r in par.runs:
                if r.font.size and 7.5<=r.font.size.pt<=8.5: r.font.size=Pt(9.5)
# Acteurs (slide 13) : corps 10 -> 11.5
for sh in p.slides[12].shapes:
    if sh.has_text_frame:
        for par in sh.text_frame.paragraphs:
            for r in par.runs:
                if r.font.size and r.font.size.pt==10: r.font.size=Pt(11.5)
p.save(OUT); print('ok')
