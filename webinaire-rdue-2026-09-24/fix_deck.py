import copy, re
from pptx import Presentation
from pptx.opc.constants import RELATIONSHIP_TYPE as RT

SRC='AtelierRDUECCI.pptx'; OUT='AtelierRDUECCI_revu_2026-09-16.pptx'
p=Presentation(SRC)
S=lambda n: p.slides[n-1]
def sh(n,i): return list(S(n).shapes)[i]
def setp(shape, pi, text, run=0):
    par=shape.text_frame.paragraphs[pi]; runs=par.runs
    runs[run].text=text
    for r in runs[run+1:]: r.text=''
def set1(shape, text): setp(shape,0,text)
def add_note(n, txt):
    ns=S(n).notes_slide.notes_text_frame
    ns.text=(ns.text.strip()+' '+txt).strip()

# ---- S6 Le RDUE en bref
setp(sh(6,6),1,"Règlement (UE) 2023/1115 du 31 mai 2023, modifié par le règlement (UE) 2025/2650 du 19 décembre 2025 (JOUE du 23/12/2025, en vigueur le 26/12/2025). Il abroge le règlement bois RBUE (995/2010) à la date d'application.")
setp(sh(6,14),0,"Grandes & moyennes entreprises (micro & petites : 30/06/2027).")

# ---- S7 calendrier (6 jalons)
set1(sh(7,5),"Reporté d'un an par le règlement (UE) 2025/2650 — calendrier confirmé par la revue de simplification du 4 mai 2026.")
r=sh(7,9).text_frame.paragraphs[0].runs
r[0].text="REPORT D'UN AN  "; r[1].text="— l'échéance initiale (2025) est décalée à 2026/2027. Le 4 mai 2026, la Commission a confirmé qu'elle ne rouvrira pas le texte."
jalons=[("26/12/2025","Modificatif en vigueur","règl. 2025/2650 : report + simplifications"),
        ("fin juin 2026","Système d'info. rouvert","dépôt des DDR (ex-TRACES)"),
        ("13/07/2026","Acte délégué annexe I","+ règl. exéc. 2026/1565 (SI)"),
        ("30/12/2026","Socle","grandes & moyennes ent."),
        ("30/06/2027","Micro & petites ent.","fin du report TPE/PE"),
        ("30/12/2027","Nouveaux produits","ajouts annexe I")]
for (a,b,c),base in zip(jalons,(13,18,23,28,33,38)):
    set1(sh(7,base),a); set1(sh(7,base+1),b); set1(sh(7,base+2),c)
set1(sh(7,41),"Reporté n'est pas supprimé — la date butoir du 31/12/2020 est inchangée ; la traçabilité jusqu'à la parcelle se prépare dès maintenant.")
add_note(7,"Acte délégué annexe I : fin de la période d'examen le 13/09/2026, publication au JOUE attendue mi-septembre — vérifier la veille du webinaire.")

# ---- S9 matières
set1(sh(9,30),"Bœuf"); set1(sh(9,31),"NC 0102·0201 (cuir : exclu*)")
set1(sh(9,37),"Dérivés : papier, charbon de bois, pneus, chocolat… vous font entrer dans le champ. *Cuirs et peaux : sortis par l'acte délégué du 13/07/2026.")

# ---- S12 acteurs
set1(sh(12,21),"art. 2 · règl. 2025/2650")
b=sh(12,23); setp(b,0,"Transforme un produit déjà mis sur le marché UE."); setp(b,1,"Pas de diligence, pas de DDR à déposer.")
setp(b,2,"1er aval : collecte et conserve le n° de DDR amont (5 ans)."); setp(b,3,"Avals suivants : infos fournisseur seulement.")
c=sh(12,32); setp(c,0,"Met à disposition sans transformer."); setp(c,1,"1er commerçant aval : conserve le n° de DDR.")
setp(c,2,"Suivants : nom, adresse, site du fournisseur."); setp(c,3,"Enregistrement au SI si non-PME.")
add_note(12,"Depuis 2025/2650 : seul le 1er opérateur/commerçant aval collecte le n° de DDR ; pas de retransmission plus loin.")

# ---- S14 annexe I
set1(sh(14,4),"L'évolution de l'annexe I — acte délégué du 13 juillet 2026")
set1(sh(14,5),"Exclusions dès l'entrée en vigueur (JOUE attendu mi-sept. 2026) ; ajouts applicables au 30/12/2027.")
a=sh(14,11); setp(a,0,"Café soluble (SH 2101 11 00)"); setp(a,1,"Dérivés d'huile de palme, acides oléiques, savons"); setp(a,2,"Langues de bovins congelées")
e=sh(14,17); setp(e,0,"Cuirs et peaux de bovins (SH 4101, 4104, 4107)"); setp(e,1,"Pneus rechapés, courroies, certains articles en caoutchouc")
setp(e,2,"Soja de semence · sièges d'avion et de véhicules"); setp(e,3,"(Produits imprimés : déjà sortis via le règl. 2025/2650)")
add_note(14,"Vérifier la version publiée au JOUE ; liste indicative issue de l'acte adopté le 13/07/2026.")

# ---- S18 pays
set1(sh(18,5),"Le pays de production module la diligence (benchmarking, art. 29 — liste du 22/05/2025, révision annoncée pour 2026).")

# ---- S23 SI
set1(sh(23,4),"Le système d'information RDUE & EU Login")
set1(sh(23,10),"Créer un compte EU Login, puis s'enregistrer dans le système d'information RDUE (rouvert fin juin 2026).")
set1(sh(23,14),"La DDR dans le système d'information (TRACES) — à saisir")
add_note(23,"Règl. d'exécution (UE) 2026/1565 en vigueur le 17/07/2026 : déclaration simplifiée micro/petits producteurs, API, groupage. Formations Commission en septembre 2026.")

# ---- S24 douane
set1(sh(24,5),"Le numéro de DDR se reporte dans la déclaration douanière (Delta), avant mainlevée.")
set1(sh(24,13),"PME aval (DDR amont existante) : C717 + n° de DDR du fournisseur — à confirmer sur la fiche douane 2026.")
set1(sh(24,32),"L'opérateur transmet le n° de DDR et le code de vérification à son seul client aval immédiat ; les maillons suivants ne conservent que les infos fournisseur.")
set1(sh(24,33),"Autres codes : Y142 (activité non commerciale) · 99EU9999999999 = n° conventionnel export/réimport, produits fabriqués entre le 29/06/2023 et l'application.")
add_note(24,"Une DDR citée en douane ne peut plus être modifiée ni retirée : vérifier quantité, SH et géolocalisation avant.")

# ---- S25 FAQ
set1(sh(25,5),"Ce que la 5ᵉ édition de la FAQ (avril 2026) confirme ou clarifie — n° de question à contrôler sur la v5.")
set1(sh(25,12),"FAQ 3.17 à 3.19")
set1(sh(25,31),"Dans le système d'information, un même compte porte opérateur / aval / commerçant / mandataire.")

# ---- S28 / S29 sanctions
set1(sh(28,20),"Amende (plafond ≥ 4 % du CA UE), confiscation, exclusion.")
set1(sh(29,10),"Plafond fixé par chaque État membre à au moins 4 % du CA annuel UE.")
set1(sh(29,20),"Exclusion temporaire (max. 12 mois) des marchés et financements publics.")
set1(sh(29,25),"Poursuites possibles (directive (UE) 2024/1203).")

# ---- S30 leviers
set1(sh(30,10),"S'appuyer sur la DDR de l'amont plutôt que reconstituer la traçabilité.")
set1(sh(30,30),"Déclaration unique pour les micro/petits producteurs de pays à risque faible.")

# ---- Quiz final : copie de la slide 13
src=S(13); new=p.slides.add_slide(p.slide_layouts[0])
for shp in list(new.shapes): shp._element.getparent().remove(shp._element)
for shp in src.shapes:
    el=copy.deepcopy(shp._element)
    if shp.shape_type==13:
        rid=new.part.relate_to(shp.part.related_part(shp._element.blip_rId), RT.IMAGE)
        for blip in el.iter('{http://schemas.openxmlformats.org/drawingml/2006/main}blip'):
            blip.set('{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed', rid)
    new.shapes._spTree.append(el)
q=list(new.shapes)
q[4].text_frame.paragraphs[0].runs[0].text="Vérifions ensemble : 3 questions"
q[5].text_frame.paragraphs[0].runs[0].text="Sondage dans le chat, puis correction."
for lbl,qt,ans,(a,b,c) in [("Q1","Un produit de l'annexe I fabriqué le 12 mars 2023 : quel code portez-vous en douane ?","Y132 — fabriqué avant le 29/06/2023",(9,10,12)),
                            ("Q2","Le RDUE s'applique-t-il aux exportations depuis l'Union européenne ?","Oui — mise sur le marché ET exportation",(16,17,19)),
                            ("Q3","« Zéro déforestation » : aucune déforestation de la parcelle après quelle date ?","Le 31 décembre 2020",(23,24,26))]:
    q[a].text_frame.paragraphs[0].runs[0].text=lbl; q[b].text_frame.paragraphs[0].runs[0].text=qt; q[c].text_frame.paragraphs[0].runs[0].text=ans
new.notes_slide.notes_text_frame.text="Quiz final : masquer les réponses (animation) ou les révéler après le sondage."

# ---- Réordonnancement
order=[1,2,3,4,5,6,7, 8,9,14,10,11,12,13,25, 15,16,17,18,19, 22,23,24,21, 27,28,29,30,26,20,31,32,34,33]
assert sorted(order)==list(range(1,35))
lst=p.slides._sldIdLst; ids=list(lst)
for el in ids: lst.remove(el)
for n in order: lst.append(ids[n-1])

# ---- Renumérotation des pieds de page
tot=len(p.slides)
for i,s in enumerate(p.slides,1):
    for shp in s.shapes:
        if shp.has_text_frame and re.fullmatch(r'\d+ / 33', shp.text_frame.text.strip()):
            shp.text_frame.paragraphs[0].runs[0].text=f"{i} / {tot}"
p.save(OUT); print("saved",OUT,tot,"slides")
